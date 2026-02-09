# -*- coding: utf-8 -*-
"""
Import Dialog
=============
ديالوج استيراد البيانات من Excel مع التحقق بالذكاء الاصطناعي

Features:
- Select Excel file
- Preview data with validation status
- AI-powered validation via Ollama (if available)
- Smart duplicate detection
- Error reporting with row/column details
- Import valid records only
- Dark/Light theme support
"""

from typing import Optional, List, Dict, Tuple
import os

from PyQt5.QtWidgets import (
    QDialog, QVBoxLayout, QHBoxLayout, QLabel,
    QPushButton, QFileDialog, QFrame, QTableWidget,
    QTableWidgetItem, QProgressBar, QHeaderView,
    QAbstractItemView, QCheckBox
)
from PyQt5.QtCore import Qt, QThread, pyqtSignal
from PyQt5.QtGui import QFont, QColor

from core.database.queries import select_all, insert_returning_id, get_count
from core.themes import get_current_theme
from core.logging import app_logger
from ui.components.notifications import toast_error, toast_warning, toast_info


class AIValidationWorker(QThread):
    """Worker thread for AI-powered validation."""

    finished = pyqtSignal(str)  # AI analysis result
    error = pyqtSignal(str)

    def __init__(self, data_preview: str, entity_name: str):
        super().__init__()
        self._data_preview = data_preview
        self._entity_name = entity_name

    def run(self):
        """Run AI validation in background."""
        try:
            from core.ai.ollama_client import get_ollama_client

            client = get_ollama_client()
            if not client.is_available():
                self.finished.emit("")
                return

            prompt = (
                f"أنت مساعد ذكي لنظام إدارة بيانات الموظفين INTEGRA.\n"
                f"راجع البيانات التالية المستوردة لجدول '{self._entity_name}':\n\n"
                f"{self._data_preview}\n\n"
                f"المطلوب:\n"
                f"1. هل توجد أخطاء إملائية واضحة في الأسماء العربية أو الإنجليزية؟\n"
                f"2. هل توجد بيانات مكررة أو متشابهة جداً؟\n"
                f"3. هل توجد بيانات غير منطقية أو غير متسقة؟\n"
                f"4. اقتراحات لتحسين البيانات.\n\n"
                f"أجب بشكل مختصر ومنظم باللغة العربية."
            )

            response = client.chat(
                message=prompt,
                system="أنت مدقق بيانات محترف. أجب بشكل مختصر ودقيق.",
                temperature=0.3
            )

            self.finished.emit(response or "")

        except Exception as e:
            app_logger.error(f"AI validation error: {e}", exc_info=True)
            self.error.emit(str(e))


class ImportDialog(QDialog):
    """
    Professional Excel import dialog with AI validation.

    Workflow:
    1. Select Excel file
    2. Preview data in table
    3. Run validation (basic + AI)
    4. Show errors/warnings
    5. Import valid records
    """

    def __init__(self, entity_key: str, parent=None):
        super().__init__(parent)

        from .master_data_window import ENTITY_CONFIGS
        if entity_key not in ENTITY_CONFIGS:
            raise ValueError(f"Unknown entity: {entity_key}")

        self._entity_key = entity_key
        self._config = ENTITY_CONFIGS[entity_key]
        self._raw_data: List[Dict] = []
        self._validation_errors: List[Dict] = []
        self._ai_worker: Optional[AIValidationWorker] = None
        self._import_success = False

        self._setup_ui()
        self._apply_theme()

    def _setup_ui(self):
        """Setup dialog UI."""
        icon = self._config['icon']
        title_ar = self._config['title_ar']
        self.setWindowTitle(f"📥 استيراد {title_ar} من Excel")
        self.setMinimumSize(800, 600)
        self.setModal(True)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 20, 20, 20)
        layout.setSpacing(15)

        # Header
        header = QLabel(f"📥 استيراد بيانات {title_ar}")
        header.setFont(QFont("Cairo", 16, QFont.Bold))
        header.setAlignment(Qt.AlignCenter)
        header.setObjectName("importHeader")
        layout.addWidget(header)

        # Instructions
        fields_text = " | ".join(
            f.get('label', f.get('key', ''))
            for f in self._config['fields']
        )
        instructions = QLabel(
            f"الأعمدة المطلوبة في ملف Excel: {fields_text}\n"
            f"يمكن أن يكون ترتيب الأعمدة في السطر الأول كعناوين."
        )
        instructions.setFont(QFont("Cairo", 11))
        instructions.setObjectName("importInstructions")
        instructions.setWordWrap(True)
        instructions.setAlignment(Qt.AlignCenter)
        layout.addWidget(instructions)

        # File selection
        file_frame = QFrame()
        file_frame.setObjectName("fileFrame")
        file_layout = QHBoxLayout(file_frame)
        file_layout.setContentsMargins(15, 10, 15, 10)

        self._file_label = QLabel("لم يتم اختيار ملف")
        self._file_label.setFont(QFont("Cairo", 11))
        self._file_label.setObjectName("fileLabel")
        file_layout.addWidget(self._file_label, 1)

        browse_btn = QPushButton("📂 اختيار ملف Excel")
        browse_btn.setFont(QFont("Cairo", 11))
        browse_btn.setCursor(Qt.PointingHandCursor)
        browse_btn.setProperty("primary", True)
        browse_btn.clicked.connect(self._browse_file)
        file_layout.addWidget(browse_btn)

        layout.addWidget(file_frame)

        # Preview table
        self._preview_table = QTableWidget()
        self._preview_table.setObjectName("previewTable")
        self._preview_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self._preview_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self._preview_table.horizontalHeader().setStretchLastSection(True)
        self._preview_table.setAlternatingRowColors(True)
        layout.addWidget(self._preview_table, 1)

        # AI Analysis label
        self._ai_label = QLabel("")
        self._ai_label.setFont(QFont("Cairo", 11))
        self._ai_label.setObjectName("aiLabel")
        self._ai_label.setWordWrap(True)
        self._ai_label.setVisible(False)
        layout.addWidget(self._ai_label)

        # Validation summary
        self._validation_label = QLabel("")
        self._validation_label.setFont(QFont("Cairo", 12, QFont.Bold))
        self._validation_label.setObjectName("validationLabel")
        self._validation_label.setAlignment(Qt.AlignCenter)
        self._validation_label.setVisible(False)
        layout.addWidget(self._validation_label)

        # Progress bar
        self._progress = QProgressBar()
        self._progress.setObjectName("importProgress")
        self._progress.setVisible(False)
        layout.addWidget(self._progress)

        # Skip duplicates checkbox
        self._skip_duplicates = QCheckBox("تخطي السجلات المكررة تلقائياً")
        self._skip_duplicates.setFont(QFont("Cairo", 11))
        self._skip_duplicates.setChecked(True)
        self._skip_duplicates.setObjectName("skipDuplicates")
        layout.addWidget(self._skip_duplicates)

        # Buttons
        buttons_layout = QHBoxLayout()
        buttons_layout.addStretch()

        cancel_btn = QPushButton("❌ إلغاء")
        cancel_btn.setFont(QFont("Cairo", 12))
        cancel_btn.setMinimumHeight(44)
        cancel_btn.setCursor(Qt.PointingHandCursor)
        cancel_btn.clicked.connect(self.reject)
        buttons_layout.addWidget(cancel_btn)

        self._validate_btn = QPushButton("🤖 تحقق بالذكاء الاصطناعي")
        self._validate_btn.setFont(QFont("Cairo", 12))
        self._validate_btn.setMinimumHeight(44)
        self._validate_btn.setCursor(Qt.PointingHandCursor)
        self._validate_btn.setProperty("primary", True)
        self._validate_btn.setEnabled(False)
        self._validate_btn.clicked.connect(self._run_ai_validation)
        buttons_layout.addWidget(self._validate_btn)

        self._import_btn = QPushButton("📥 استيراد البيانات")
        self._import_btn.setFont(QFont("Cairo", 12, QFont.Bold))
        self._import_btn.setMinimumHeight(44)
        self._import_btn.setCursor(Qt.PointingHandCursor)
        self._import_btn.setProperty("buttonColor", "success")
        self._import_btn.setEnabled(False)
        self._import_btn.clicked.connect(self._do_import)
        buttons_layout.addWidget(self._import_btn)

        buttons_layout.addStretch()
        layout.addLayout(buttons_layout)

    # ═══════════════════════════════════════════════════════════════
    # File & Data Loading
    # ═══════════════════════════════════════════════════════════════

    def _browse_file(self):
        """Open file browser to select Excel file."""
        filepath, _ = QFileDialog.getOpenFileName(
            self,
            "اختيار ملف Excel",
            "",
            "Excel Files (*.xlsx *.xls);;All Files (*)"
        )
        if not filepath:
            return

        self._file_label.setText(os.path.basename(filepath))
        self._load_excel(filepath)

    def _load_excel(self, filepath: str):
        """Load Excel file and preview data."""
        try:
            import openpyxl
        except ImportError:
            toast_error(self, "خطأ", "مكتبة openpyxl غير مثبتة!")
            return

        try:
            wb = openpyxl.load_workbook(filepath, read_only=True)
            ws = wb.active

            if ws is None:
                toast_error(self, "خطأ", "الملف لا يحتوي على بيانات!")
                return

            rows = list(ws.iter_rows(values_only=True))
            wb.close()

            if len(rows) < 2:
                toast_warning(self, "تنبيه", "الملف يحتوي على أقل من سطرين (عناوين + بيانات)")
                return

            # Extract field keys from config
            field_keys = [f['key'] for f in self._config['fields']]
            field_labels = [f['label'] for f in self._config['fields']]

            # First row as headers (skip it)
            header_row = rows[0]

            # Map columns: try to match headers to field labels/keys
            col_mapping = self._auto_map_columns(header_row, field_keys, field_labels)

            # Parse data rows
            self._raw_data = []
            for row_idx, row in enumerate(rows[1:], start=2):
                record = {}
                for field_idx, field_key in enumerate(field_keys):
                    col_idx = col_mapping.get(field_key)
                    if col_idx is not None and col_idx < len(row):
                        value = row[col_idx]
                        record[field_key] = str(value).strip() if value is not None else ''
                    else:
                        record[field_key] = ''
                record['_row'] = row_idx
                self._raw_data.append(record)

            # Run basic validation
            self._run_basic_validation()

            # Show preview
            self._show_preview()

            # Enable buttons
            self._validate_btn.setEnabled(True)
            self._import_btn.setEnabled(True)

        except Exception as e:
            app_logger.error(f"Excel load error: {e}", exc_info=True)
            toast_error(self, "خطأ", f"فشل قراءة الملف: {e}")

    def _auto_map_columns(
        self,
        header_row: tuple,
        field_keys: List[str],
        field_labels: List[str]
    ) -> Dict[str, int]:
        """
        Auto-map Excel columns to entity fields.
        Tries to match by label (Arabic), then by key (English).
        Falls back to positional mapping.
        """
        mapping = {}
        headers = [str(h).strip().lower() if h else '' for h in header_row]

        for i, (key, label) in enumerate(zip(field_keys, field_labels)):
            # Try exact label match
            label_lower = label.lower()
            key_lower = key.lower()

            matched = False
            for col_idx, header in enumerate(headers):
                if header == label_lower or header == key_lower:
                    mapping[key] = col_idx
                    matched = True
                    break

            # Try partial match
            if not matched:
                for col_idx, header in enumerate(headers):
                    if label_lower in header or key_lower in header:
                        if col_idx not in mapping.values():
                            mapping[key] = col_idx
                            matched = True
                            break

            # Fallback: positional
            if not matched and i < len(header_row):
                mapping[key] = i

        return mapping

    # ═══════════════════════════════════════════════════════════════
    # Validation
    # ═══════════════════════════════════════════════════════════════

    def _run_basic_validation(self):
        """Run basic validation on loaded data."""
        self._validation_errors = []

        # Get existing records for duplicate detection
        existing = self._get_existing_values()

        for record in self._raw_data:
            row_num = record.get('_row', '?')
            errors = []

            for field_def in self._config['fields']:
                key = field_def['key']
                value = record.get(key, '')
                is_required = field_def.get('required', False)

                # Required check
                if is_required and not value:
                    errors.append(f"حقل '{field_def['label']}' مطلوب")

                # Max length
                max_len = field_def.get('max_length')
                if value and max_len and len(value) > max_len:
                    errors.append(
                        f"حقل '{field_def['label']}' تجاوز الحد ({len(value)}/{max_len})"
                    )

            # Duplicate in existing DB
            name_ar = record.get('name_ar', '')
            if name_ar and name_ar.lower() in existing.get('name_ar', set()):
                errors.append(f"'{name_ar}' موجود بالفعل في قاعدة البيانات")

            # Duplicate within file
            name_ar_list = [
                r.get('name_ar', '').lower()
                for r in self._raw_data
                if r.get('_row') != row_num
            ]
            if name_ar and name_ar.lower() in name_ar_list:
                errors.append(f"'{name_ar}' مكرر داخل الملف")

            record['_errors'] = errors
            record['_valid'] = len(errors) == 0

    def _get_existing_values(self) -> Dict[str, set]:
        """Get existing values from database for duplicate detection."""
        existing: Dict[str, set] = {'name_ar': set(), 'name_en': set()}
        try:
            table = self._config['table']
            columns, rows = select_all(f"SELECT name_ar, name_en FROM {table}")
            if rows:
                for row in rows:
                    if row[0]:
                        existing['name_ar'].add(str(row[0]).lower())
                    if row[1]:
                        existing['name_en'].add(str(row[1]).lower())
        except Exception as e:
            app_logger.error(f"Error fetching existing data: {e}", exc_info=True)
        return existing

    def _run_ai_validation(self):
        """Run AI-powered validation using Ollama."""
        if not self._raw_data:
            toast_warning(self, "تنبيه", "لا توجد بيانات للتحقق!")
            return

        self._validate_btn.setEnabled(False)
        self._validate_btn.setText("🤖 جاري التحليل...")

        # Prepare data preview for AI
        preview_lines = []
        field_keys = [f['key'] for f in self._config['fields']]
        for record in self._raw_data[:50]:  # Limit to 50 rows
            parts = [f"{k}: {record.get(k, '')}" for k in field_keys]
            preview_lines.append(" | ".join(parts))

        data_preview = "\n".join(preview_lines)

        self._ai_worker = AIValidationWorker(data_preview, self._config['title_ar'])
        self._ai_worker.finished.connect(self._on_ai_finished)
        self._ai_worker.error.connect(self._on_ai_error)
        self._ai_worker.start()

    def _on_ai_finished(self, analysis: str):
        """Handle AI validation result."""
        self._validate_btn.setEnabled(True)
        self._validate_btn.setText("🤖 تحقق بالذكاء الاصطناعي")

        if analysis:
            self._ai_label.setText(f"🤖 تحليل الذكاء الاصطناعي:\n{analysis}")
            self._ai_label.setVisible(True)
        else:
            self._ai_label.setText("⚠️ الذكاء الاصطناعي غير متاح حالياً - تم التحقق الأساسي فقط")
            self._ai_label.setVisible(True)

    def _on_ai_error(self, error_msg: str):
        """Handle AI validation error."""
        self._validate_btn.setEnabled(True)
        self._validate_btn.setText("🤖 تحقق بالذكاء الاصطناعي")
        self._ai_label.setText(f"⚠️ خطأ في الذكاء الاصطناعي: {error_msg}")
        self._ai_label.setVisible(True)

    # ═══════════════════════════════════════════════════════════════
    # Preview
    # ═══════════════════════════════════════════════════════════════

    def _show_preview(self):
        """Show data preview in table."""
        if not self._raw_data:
            return

        field_keys = [f['key'] for f in self._config['fields']]
        field_labels = [f['label'] for f in self._config['fields']]

        # Setup columns: Status + fields + Errors
        col_headers = ['الحالة'] + field_labels + ['الأخطاء']
        self._preview_table.setColumnCount(len(col_headers))
        self._preview_table.setHorizontalHeaderLabels(col_headers)
        self._preview_table.setRowCount(len(self._raw_data))

        valid_count = 0
        error_count = 0

        for row_idx, record in enumerate(self._raw_data):
            is_valid = record.get('_valid', True)
            errors = record.get('_errors', [])

            # Status column
            status_item = QTableWidgetItem("✅" if is_valid else "❌")
            status_item.setTextAlignment(Qt.AlignCenter)
            self._preview_table.setItem(row_idx, 0, status_item)

            # Data columns
            for col_idx, key in enumerate(field_keys):
                value = record.get(key, '')
                item = QTableWidgetItem(str(value))
                item.setTextAlignment(Qt.AlignCenter)
                if not is_valid:
                    item.setForeground(QColor('#f87171'))
                self._preview_table.setItem(row_idx, col_idx + 1, item)

            # Errors column
            error_text = ' | '.join(errors) if errors else ''
            error_item = QTableWidgetItem(error_text)
            error_item.setForeground(QColor('#f87171'))
            self._preview_table.setItem(row_idx, len(field_keys) + 1, error_item)

            if is_valid:
                valid_count += 1
            else:
                error_count += 1

        # Resize columns
        self._preview_table.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeToContents
        )

        # Show summary
        total = len(self._raw_data)
        self._validation_label.setText(
            f"المجموع: {total} | صالح: {valid_count} ✅ | "
            f"أخطاء: {error_count} ❌"
        )
        self._validation_label.setVisible(True)

    # ═══════════════════════════════════════════════════════════════
    # Import
    # ═══════════════════════════════════════════════════════════════

    def _do_import(self):
        """Import valid records into database."""
        if not self._raw_data:
            toast_warning(self, "تنبيه", "لا توجد بيانات للاستيراد!")
            return

        skip_duplicates = self._skip_duplicates.isChecked()

        # Filter valid records
        records_to_import = [
            r for r in self._raw_data
            if r.get('_valid', False) or (skip_duplicates and self._is_only_duplicate_error(r))
        ]

        if not records_to_import:
            toast_warning(self, "تنبيه", "لا توجد سجلات صالحة للاستيراد!")
            return

        self._import_btn.setEnabled(False)
        self._progress.setVisible(True)
        self._progress.setMaximum(len(records_to_import))

        success_count = 0
        skip_count = 0
        fail_count = 0
        table = self._config['table']
        field_keys = [f['key'] for f in self._config['fields']]

        for idx, record in enumerate(records_to_import):
            try:
                # Skip if already exists and skip_duplicates is on
                name_ar = record.get('name_ar', '')
                if skip_duplicates and name_ar:
                    count = get_count(
                        f"SELECT COUNT(*) FROM {table} WHERE name_ar = %s",
                        (name_ar,)
                    )
                    if count and count > 0:
                        skip_count += 1
                        self._progress.setValue(idx + 1)
                        continue

                # Build insert query
                fields = []
                values = []
                for key in field_keys:
                    value = record.get(key, '')
                    if value:
                        fields.append(key)
                        values.append(value)

                if fields:
                    columns_str = ', '.join(fields)
                    placeholders = ', '.join(['%s'] * len(fields))
                    query = f"INSERT INTO {table} ({columns_str}) VALUES ({placeholders}) RETURNING id"
                    new_id = insert_returning_id(query, tuple(values))
                    if new_id:
                        success_count += 1
                    else:
                        fail_count += 1

            except Exception as e:
                app_logger.error(f"Import row error: {e}", exc_info=True)
                fail_count += 1

            self._progress.setValue(idx + 1)

        # Show result
        result_msg = f"تم الاستيراد: {success_count} ✅"
        if skip_count > 0:
            result_msg += f" | تم تخطي: {skip_count} مكرر"
        if fail_count > 0:
            result_msg += f" | فشل: {fail_count} ❌"

        self._validation_label.setText(result_msg)

        if success_count > 0:
            self._import_success = True
            toast_info(self, "نتيجة الاستيراد", result_msg)
            self.accept()
        else:
            self._import_btn.setEnabled(True)
            toast_warning(self, "تنبيه", "لم يتم استيراد أي سجل!")

    def _is_only_duplicate_error(self, record: Dict) -> bool:
        """Check if the only errors are duplicate-related."""
        errors = record.get('_errors', [])
        if not errors:
            return True
        return all('موجود بالفعل' in e or 'مكرر' in e for e in errors)

    def exec_(self) -> int:
        """Override exec_ to return proper result."""
        result = super().exec_()
        if self._import_success:
            return QDialog.Accepted
        return result

    # ═══════════════════════════════════════════════════════════════
    # Theme
    # ═══════════════════════════════════════════════════════════════

    def _apply_theme(self):
        """Apply current theme."""
        theme = get_current_theme()

        if theme == 'dark':
            self.setStyleSheet("""
                QDialog { background-color: #1e293b; }

                QLabel { color: #f1f5f9; background: transparent; }
                QLabel#importHeader { color: #38bdf8; }
                QLabel#importInstructions { color: #94a3b8; }
                QLabel#fileLabel { color: #94a3b8; }
                QLabel#aiLabel {
                    color: #a5f3fc; background-color: #0f172a;
                    border: 1px solid #164e63; border-radius: 8px;
                    padding: 12px; font-size: 12px;
                }
                QLabel#validationLabel { color: #38bdf8; }

                QFrame#fileFrame {
                    background-color: #0f172a;
                    border: 1px solid #334155;
                    border-radius: 8px;
                }

                QTableWidget#previewTable {
                    background-color: #0f172a;
                    color: #f1f5f9;
                    border: 1px solid #334155;
                    border-radius: 8px;
                    gridline-color: #334155;
                    selection-background-color: #1e40af;
                }
                QTableWidget#previewTable QHeaderView::section {
                    background-color: #1e293b;
                    color: #94a3b8;
                    border: 1px solid #334155;
                    padding: 6px;
                    font-weight: bold;
                    font-family: Cairo;
                }
                QTableWidget#previewTable::item:alternate {
                    background-color: #1e293b;
                }

                QCheckBox#skipDuplicates { color: #94a3b8; spacing: 8px; }

                QProgressBar#importProgress {
                    background-color: #0f172a; border: none;
                    border-radius: 4px; height: 8px;
                }
                QProgressBar#importProgress::chunk {
                    background-color: #10b981; border-radius: 4px;
                }

                QPushButton {
                    background-color: #334155; color: #f1f5f9;
                    border: none; border-radius: 8px;
                    padding: 10px 20px; font-weight: bold;
                }
                QPushButton:hover { background-color: #475569; }
                QPushButton:disabled { background-color: #1e293b; color: #475569; }
                QPushButton[primary="true"] { background-color: #2563eb; }
                QPushButton[primary="true"]:hover { background-color: #1d4ed8; }
                QPushButton[buttonColor="success"] { background-color: #10b981; }
                QPushButton[buttonColor="success"]:hover { background-color: #059669; }
            """)
        else:
            self.setStyleSheet("""
                QDialog { background-color: #ffffff; }

                QLabel { color: #1e293b; background: transparent; }
                QLabel#importHeader { color: #0891b2; }
                QLabel#importInstructions { color: #64748b; }
                QLabel#fileLabel { color: #64748b; }
                QLabel#aiLabel {
                    color: #155e75; background-color: #ecfeff;
                    border: 1px solid #a5f3fc; border-radius: 8px;
                    padding: 12px; font-size: 12px;
                }
                QLabel#validationLabel { color: #0891b2; }

                QFrame#fileFrame {
                    background-color: #f8fafc;
                    border: 1px solid #e2e8f0;
                    border-radius: 8px;
                }

                QTableWidget#previewTable {
                    background-color: #ffffff;
                    color: #1e293b;
                    border: 1px solid #e2e8f0;
                    border-radius: 8px;
                    gridline-color: #e2e8f0;
                    selection-background-color: #dbeafe;
                }
                QTableWidget#previewTable QHeaderView::section {
                    background-color: #f1f5f9;
                    color: #475569;
                    border: 1px solid #e2e8f0;
                    padding: 6px;
                    font-weight: bold;
                    font-family: Cairo;
                }

                QCheckBox#skipDuplicates { color: #64748b; spacing: 8px; }

                QProgressBar#importProgress {
                    background-color: #f1f5f9; border: none;
                    border-radius: 4px; height: 8px;
                }
                QProgressBar#importProgress::chunk {
                    background-color: #10b981; border-radius: 4px;
                }

                QPushButton {
                    background-color: #e2e8f0; color: #1e293b;
                    border: none; border-radius: 8px;
                    padding: 10px 20px; font-weight: bold;
                }
                QPushButton:hover { background-color: #cbd5e1; }
                QPushButton:disabled { background-color: #f1f5f9; color: #94a3b8; }
                QPushButton[primary="true"] { background-color: #2563eb; color: #ffffff; }
                QPushButton[primary="true"]:hover { background-color: #1d4ed8; }
                QPushButton[buttonColor="success"] { background-color: #10b981; color: #ffffff; }
                QPushButton[buttonColor="success"]:hover { background-color: #059669; }
            """)
